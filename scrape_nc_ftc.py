#!/usr/bin/env python3
"""
Scrape NC FTC 2025 DECODE season advancement points from all 13 qualifier events
and produce a summary Excel spreadsheet.
"""

import re
import sys
import time

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Event definitions
# ---------------------------------------------------------------------------
EVENTS = [
    ("USNCRAQ",  "Cardinal Gibbons",  "1/17"),
    ("USNCROQ",  "Thales Academy",    "1/17"),
    ("USNCCOQ",  "TMSA Charlotte",    "1/17"),
    ("USNCSAQ",  "Ascend Leadership", "2/7"),
    ("USNCSHQ",  "Pinnacle Classical","2/7"),
    ("USNCRAQ2", "SE Raleigh 2",      "2/7"),
    ("USNCSAQ2", "Ascend Leadership 2","2/8"),
    ("USNCSHQ2", "Pinnacle Classical 2","2/8"),
    ("USNCASQ",  "Asheville School",  "2/14"),
    ("USNCGRQ",  "CM Eppes MS",       "2/14"),
    ("USNCWSQ2", "Salem Academy 2",   "2/15"),
    ("USNCGRQ2", "SE Guilford 2",     "2/15"),
    ("USNCRAQ3", "SE Raleigh 3",      "2/15"),
]

BASE_URL = "https://ftc-events.firstinspires.org/2025/{code}/advancementpoints"


def scrape_event(code: str, name: str) -> list[dict] | None:
    """Scrape a single event's advancement points table.

    Returns a list of dicts with keys:
        team_number, team_name, total_pts, advanced
    or None if the page has no data.
    """
    url = BASE_URL.format(code=code)
    try:
        resp = requests.get(url, timeout=30)
        resp.raise_for_status()
    except requests.RequestException as exc:
        print(f"  WARNING: Could not fetch {code} ({name}): {exc}")
        return None

    soup = BeautifulSoup(resp.text, "html.parser")
    table = soup.find("table", id="advancementPointsTable")
    if table is None:
        print(f"  WARNING: No advancement table found for {code} ({name})")
        return None

    tbody = table.find("tbody")
    if tbody is None:
        print(f"  WARNING: No tbody found for {code} ({name})")
        return None

    rows = tbody.find_all("tr")
    if not rows:
        print(f"  WARNING: Empty table for {code} ({name})")
        return None

    results = []
    for tr in rows:
        cells = tr.find_all("td")
        if len(cells) < 9:
            continue

        # Team number: inside <a> tag in second cell
        team_num_tag = cells[1].find("a")
        if team_num_tag:
            team_number = team_num_tag.get_text(strip=True)
        else:
            team_number = cells[1].get_text(strip=True)

        team_name = cells[2].get_text(strip=True)

        # Total Pts (may be wrapped in <strong>)
        total_pts_text = cells[3].get_text(strip=True)
        try:
            total_pts = int(total_pts_text)
        except ValueError:
            total_pts = 0

        # # Advanced column (index 8)
        adv_text = cells[8].get_text(strip=True)
        advanced = adv_text not in ("-", "", "0")

        results.append({
            "team_number": int(team_number),
            "team_name": team_name,
            "total_pts": total_pts,
            "advanced": advanced,
        })

    return results


def build_spreadsheet(all_data: dict, output_path: str):
    """Build the Excel spreadsheet from scraped data.

    all_data: dict mapping event_code -> list of team dicts (or None)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Advancement Points"

    # Collect all teams across events
    teams: dict[int, str] = {}  # team_number -> team_name
    for code, _, _ in EVENTS:
        event_data = all_data.get(code)
        if event_data is None:
            continue
        for row in event_data:
            num = row["team_number"]
            if num not in teams:
                teams[num] = row["team_name"]

    # Build lookup: (event_code, team_number) -> {total_pts, advanced}
    lookup: dict[tuple[str, int], dict] = {}
    for code, _, _ in EVENTS:
        event_data = all_data.get(code)
        if event_data is None:
            continue
        for row in event_data:
            lookup[(code, row["team_number"])] = {
                "total_pts": row["total_pts"],
                "advanced": row["advanced"],
            }

    # Calculate totals for sorting
    team_totals = {}
    for num in teams:
        total = 0
        count = 0
        for code, _, _ in EVENTS:
            key = (code, num)
            if key in lookup:
                total += lookup[key]["total_pts"]
                count += 1
        team_totals[num] = (total, count)

    sorted_teams = sorted(teams.keys(), key=lambda n: team_totals[n][0], reverse=True)

    # Active events (ones that had data)
    active_events = [(code, name, date) for code, name, date in EVENTS if all_data.get(code) is not None]

    # ----- Header row -----
    headers = ["Team Number", "Team Name"]
    for code, name, date in active_events:
        headers.append(f"{name} {date}")
    headers += ["Events Attended", "Total Points", "Avg Points/Event"]

    bold_font = Font(bold=True)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Add hyperlinks to event column headers
    for i, (code, name, date) in enumerate(active_events):
        col_idx = 3 + i  # first event column is column 3
        cell = ws.cell(row=1, column=col_idx)
        url = BASE_URL.format(code=code)
        cell.hyperlink = url
        cell.font = Font(bold=True, color="0563C1", underline="single")

    # ----- Data rows -----
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    right_align = Alignment(horizontal="right")

    for row_idx, team_num in enumerate(sorted_teams, start=2):
        ws.cell(row=row_idx, column=1, value=team_num).alignment = right_align
        ws.cell(row=row_idx, column=2, value=teams[team_num])

        total_pts = 0
        events_attended = 0

        for event_i, (code, _, _) in enumerate(active_events):
            col_idx = 3 + event_i
            key = (code, team_num)
            if key in lookup:
                pts = lookup[key]["total_pts"]
                cell = ws.cell(row=row_idx, column=col_idx, value=pts)
                cell.alignment = right_align
                if lookup[key]["advanced"]:
                    cell.fill = green_fill
                total_pts += pts
                events_attended += 1
            # else leave blank

        num_event_cols = len(active_events)
        # Events Attended
        ea_col = 3 + num_event_cols
        ws.cell(row=row_idx, column=ea_col, value=events_attended).alignment = right_align

        # Total Points
        tp_col = ea_col + 1
        ws.cell(row=row_idx, column=tp_col, value=total_pts).alignment = right_align

        # Avg Points/Event
        avg_col = tp_col + 1
        if events_attended > 0:
            avg = round(total_pts / events_attended, 1)
        else:
            avg = 0
        ws.cell(row=row_idx, column=avg_col, value=avg).alignment = right_align

    # ----- Formatting -----
    # Auto-fit column widths
    for col_idx in range(1, len(headers) + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=False):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        # Cap event columns at a reasonable width
        if 3 <= col_idx <= 2 + len(active_events):
            adjusted = min(max_len + 2, 20)
        else:
            adjusted = max_len + 2
        ws.column_dimensions[col_letter].width = adjusted

    # Freeze first 2 columns + header row
    ws.freeze_panes = "C2"

    wb.save(output_path)
    print(f"\nSpreadsheet saved to: {output_path}")
    print(f"Total teams: {len(sorted_teams)}")
    print(f"Events with data: {len(active_events)} / {len(EVENTS)}")


def main():
    output_path = "nc_ftc_advancement_2026.xlsx"

    all_data: dict[str, list[dict] | None] = {}
    skipped = []

    print("Scraping NC FTC 2025 DECODE advancement points...\n")

    for i, (code, name, date) in enumerate(EVENTS):
        print(f"[{i+1}/{len(EVENTS)}] {name} ({code}) - {date} ... ", end="", flush=True)
        data = scrape_event(code, name)
        if data is not None:
            print(f"{len(data)} teams")
            all_data[code] = data
        else:
            print("SKIPPED (no data)")
            skipped.append((code, name, date))
        # Be polite to the server
        if i < len(EVENTS) - 1:
            time.sleep(0.5)

    if skipped:
        print(f"\nSkipped events ({len(skipped)}):")
        for code, name, date in skipped:
            print(f"  - {name} ({code}) {date}")

    if not all_data:
        print("ERROR: No data was scraped from any event!")
        sys.exit(1)

    print(f"\nBuilding spreadsheet...")
    build_spreadsheet(all_data, output_path)


if __name__ == "__main__":
    main()
